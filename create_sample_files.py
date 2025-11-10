"""
Create sample PO and Invoice files with mismatches, duplicates, and anomalies
Generates multiple sets for testing different scenarios
"""

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import os

def create_sample_set(set_number, scenario_description):
    """Create a set of PO and Invoice files with specific issues"""
    
    # Create directories if they don't exist
    os.makedirs("sample_files", exist_ok=True)
    
    # Base data - Purchase Order
    po_data = {
        'Item': ['LAPTOP-001', 'MOUSE-002', 'KEYBOARD-003', 'MONITOR-004', 'WEBCAM-005'],
        'Description': ['Dell Laptop 15"', 'Wireless Mouse', 'Mechanical Keyboard', '27" Monitor', 'HD Webcam'],
        'Quantity': [5, 10, 8, 4, 6],
        'Unit Price': [1200.00, 25.50, 89.99, 350.00, 45.00],
        'Total': [6000.00, 255.00, 719.92, 1400.00, 270.00]
    }
    
    # Create invoice data based on scenario
    if set_number == 1:
        # Set 1: Price Mismatches
        invoice_data = {
            'Item': ['LAPTOP-001', 'MOUSE-002', 'KEYBOARD-003', 'MONITOR-004', 'WEBCAM-005'],
            'Description': ['Dell Laptop 15"', 'Wireless Mouse', 'Mechanical Keyboard', '27" Monitor', 'HD Webcam'],
            'Quantity': [5, 10, 8, 4, 6],
            'Unit Price': [1250.00, 25.50, 95.00, 350.00, 50.00],  # Price mismatches
            'Total': [6250.00, 255.00, 760.00, 1400.00, 300.00]
        }
        scenario = "Price Mismatches"
        
    elif set_number == 2:
        # Set 2: Duplicates
        invoice_data = {
            'Item': ['LAPTOP-001', 'MOUSE-002', 'MOUSE-002', 'KEYBOARD-003', 'MONITOR-004', 'WEBCAM-005', 'WEBCAM-005'],
            'Description': ['Dell Laptop 15"', 'Wireless Mouse', 'Wireless Mouse', 'Mechanical Keyboard', '27" Monitor', 'HD Webcam', 'HD Webcam'],
            'Quantity': [5, 10, 5, 8, 4, 6, 3],
            'Unit Price': [1200.00, 25.50, 25.50, 89.99, 350.00, 45.00, 45.00],
            'Total': [6000.00, 255.00, 127.50, 719.92, 1400.00, 270.00, 135.00]
        }
        scenario = "Duplicates"
        
    elif set_number == 3:
        # Set 3: Both Mismatches and Duplicates
        invoice_data = {
            'Item': ['LAPTOP-001', 'MOUSE-002', 'MOUSE-002', 'KEYBOARD-003', 'MONITOR-004', 'WEBCAM-005'],
            'Description': ['Dell Laptop 15"', 'Wireless Mouse', 'Wireless Mouse', 'Mechanical Keyboard', '27" Monitor', 'HD Webcam'],
            'Quantity': [5, 10, 3, 8, 4, 6],
            'Unit Price': [1250.00, 25.50, 25.50, 95.00, 350.00, 50.00],  # Price mismatches + duplicates
            'Total': [6250.00, 255.00, 76.50, 760.00, 1400.00, 300.00]
        }
        scenario = "Mismatches and Duplicates"
        
    elif set_number == 4:
        # Set 4: Quantity Mismatches
        invoice_data = {
            'Item': ['LAPTOP-001', 'MOUSE-002', 'KEYBOARD-003', 'MONITOR-004', 'WEBCAM-005'],
            'Description': ['Dell Laptop 15"', 'Wireless Mouse', 'Mechanical Keyboard', '27" Monitor', 'HD Webcam'],
            'Quantity': [6, 10, 8, 5, 6],  # Quantity mismatches
            'Unit Price': [1200.00, 25.50, 89.99, 350.00, 45.00],
            'Total': [7200.00, 255.00, 719.92, 1750.00, 270.00]
        }
        scenario = "Quantity Mismatches"
        
    elif set_number == 5:
        # Set 5: Multiple Issues (Mismatches, Duplicates, Anomalies)
        invoice_data = {
            'Item': ['LAPTOP-001', 'MOUSE-002', 'MOUSE-002', 'KEYBOARD-003', 'MONITOR-004', 'WEBCAM-005', 'SPEAKER-006'],
            'Description': ['Dell Laptop 15"', 'Wireless Mouse', 'Wireless Mouse', 'Mechanical Keyboard', '27" Monitor', 'HD Webcam', 'Bluetooth Speaker'],
            'Quantity': [5, 10, 2, 8, 4, 6, 3],
            'Unit Price': [1250.00, 25.50, 25.50, 95.00, 350.00, 50.00, 75.00],  # Price mismatches + new item
            'Total': [6250.00, 255.00, 51.00, 760.00, 1400.00, 300.00, 225.00]
        }
        scenario = "Multiple Issues"
        
    else:
        # Default: Simple mismatches
        invoice_data = {
            'Item': ['LAPTOP-001', 'MOUSE-002', 'KEYBOARD-003'],
            'Description': ['Dell Laptop 15"', 'Wireless Mouse', 'Mechanical Keyboard'],
            'Quantity': [5, 10, 8],
            'Unit Price': [1300.00, 30.00, 100.00],  # All prices higher
            'Total': [6500.00, 300.00, 800.00]
        }
        scenario = "Simple Mismatches"
    
    # Create DataFrames
    po_df = pd.DataFrame(po_data)
    invoice_df = pd.DataFrame(invoice_data)
    
    # Create Excel files with formatting
    po_filename = f"sample_files/sample_po_set{set_number}.xlsx"
    invoice_filename = f"sample_files/sample_invoice_set{set_number}.xlsx"
    
    # Write PO file
    with pd.ExcelWriter(po_filename, engine='openpyxl') as writer:
        po_df.to_excel(writer, sheet_name='Purchase Order', index=False)
        worksheet = writer.sheets['Purchase Order']
        
        # Format header
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")
        
        # Auto-adjust column widths
        for column in worksheet.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            worksheet.column_dimensions[column_letter].width = adjusted_width
    
    # Write Invoice file
    with pd.ExcelWriter(invoice_filename, engine='openpyxl') as writer:
        invoice_df.to_excel(writer, sheet_name='Invoice', index=False)
        worksheet = writer.sheets['Invoice']
        
        # Format header
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")
        
        # Auto-adjust column widths
        for column in worksheet.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            worksheet.column_dimensions[column_letter].width = adjusted_width
    
    print(f"✅ Created Set {set_number}: {scenario}")
    print(f"   PO: {po_filename}")
    print(f"   Invoice: {invoice_filename}")
    print()
    
    return po_filename, invoice_filename, scenario

def main():
    """Create all sample file sets"""
    print("=" * 60)
    print("Creating Sample PO and Invoice Files")
    print("=" * 60)
    print()
    
    scenarios = []
    
    # Create 5 different sets
    for i in range(1, 6):
        po_file, invoice_file, scenario = create_sample_set(i, "")
        scenarios.append({
            'set': i,
            'po_file': po_file,
            'invoice_file': invoice_file,
            'scenario': scenario
        })
    
    print("=" * 60)
    print("Summary of Created Files:")
    print("=" * 60)
    print()
    print("Set 1: Price Mismatches")
    print("  - LAPTOP-001: $1200 → $1250")
    print("  - KEYBOARD-003: $89.99 → $95.00")
    print("  - WEBCAM-005: $45.00 → $50.00")
    print()
    print("Set 2: Duplicates")
    print("  - MOUSE-002: Appears twice")
    print("  - WEBCAM-005: Appears twice")
    print()
    print("Set 3: Mismatches + Duplicates")
    print("  - Price mismatches on LAPTOP-001, KEYBOARD-003, WEBCAM-005")
    print("  - MOUSE-002 appears twice")
    print()
    print("Set 4: Quantity Mismatches")
    print("  - LAPTOP-001: Qty 5 → 6")
    print("  - MONITOR-004: Qty 4 → 5")
    print()
    print("Set 5: Multiple Issues")
    print("  - Price mismatches")
    print("  - Duplicates (MOUSE-002)")
    print("  - New item not in PO (SPEAKER-006)")
    print()
    print("=" * 60)
    print("All files created in 'sample_files' folder")
    print("=" * 60)

if __name__ == "__main__":
    main()
